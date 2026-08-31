from __future__ import annotations

import csv
from pathlib import Path
import re

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader


class KnowledgeBaseLoader:
    """Load the repository Knowledge Base into LangChain Document objects."""

    _WHITESPACE = re.compile(r"[ \t]+")
    _BLANK_LINES = re.compile(r"\n{3,}")

    def __init__(self, root: str | Path) -> None:
        path = Path(root)
        if not path.is_absolute():
            path = Path(__file__).resolve().parents[2] / path
        self.root = path
        self.project_root = Path(__file__).resolve().parents[2]

    def load(self) -> list[Document]:
        """Load supported public and internal sources with retrieval metadata."""
        if not self.root.is_dir():
            raise FileNotFoundError(f"Knowledge Base directory was not found: {self.root}")

        documents: list[Document] = []
        documents.extend(self._load_faq())
        documents.extend(self._load_html_pages())
        for path in sorted(self.root.glob("*.pdf")):
            documents.extend(self._load_pdf(path))
        for path in sorted(self.root.glob("*.docx")):
            document = self._load_docx(path)
            if document is not None:
                documents.append(document)
        return documents

    def _relative_source(self, path: Path) -> str:
        return path.relative_to(self.project_root).as_posix()

    @classmethod
    def _normalize_text(cls, value: str) -> str:
        lines = [cls._WHITESPACE.sub(" ", line).strip() for line in value.splitlines()]
        return cls._BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()

    def _load_faq(self) -> list[Document]:
        path = self.root / "FAQ.csv"
        if not path.exists():
            return []

        source = self._relative_source(path)
        documents: list[Document] = []
        with path.open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                faq_id = row["FAQ_ID"].strip()
                title = row["Question"].strip()
                content = self._normalize_text(
                    "\n".join(
                        [
                            f"Question: {title}",
                            f"Answer: {row['Answer'].strip()}",
                            f"Applies to: {row['Applies_To'].strip()}",
                            f"Related policy: {row['Related_Policy'].strip()}",
                            f"Tags: {row['Tags'].strip()}",
                        ]
                    )
                )
                documents.append(
                    Document(
                        page_content=content,
                        metadata={
                            "source_path": f"{source}#{faq_id}",
                            "source_file": source,
                            "source_type": "faq",
                            "title": title,
                            "category": row["Category"].strip(),
                            "faq_id": faq_id,
                            "related_policy": row["Related_Policy"].strip(),
                            "applies_to": row["Applies_To"].strip(),
                            "tags": row["Tags"].strip(),
                            "last_updated": row["Last_Updated"].strip(),
                            "visibility": "public",
                        },
                    )
                )
        return documents

    def _load_html_pages(self) -> list[Document]:
        html_dir = self.root / "HTML Pages"
        manifest_path = html_dir / "manifest.csv"
        manifest: dict[str, dict[str, str]] = {}
        if manifest_path.exists():
            with manifest_path.open(encoding="utf-8-sig", newline="") as handle:
                manifest = {row["filename"]: row for row in csv.DictReader(handle)}

        documents: list[Document] = []
        for path in sorted(html_dir.glob("[0-9][0-9][0-9]-*.html")):
            soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
            article = soup.find("article") or soup.find("main") or soup.body
            if article is None:
                continue
            for element in article.select("style, script, nav, .back"):
                element.decompose()
            content = self._normalize_text(article.get_text("\n"))
            if not content:
                continue

            row = manifest.get(path.name, {})
            heading = article.find("h1")
            title = row.get("title") or (
                heading.get_text(" ", strip=True) if heading else path.stem
            )
            source = self._relative_source(path)
            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source_path": source,
                        "source_file": source,
                        "source_type": "html",
                        "title": title,
                        "category": row.get("category", "Help Center"),
                        "article_id": row.get("id", ""),
                        "summary": row.get("summary", ""),
                        "visibility": "public",
                    },
                )
            )
        return documents

    def _load_pdf(self, path: Path) -> list[Document]:
        source = self._relative_source(path)
        reader = PdfReader(path)
        title = path.stem
        category = "Product Manual" if title in {"Air Purifier", "Camera", "Smart Lock"} else "Policy"
        documents: list[Document] = []
        for page_number, page in enumerate(reader.pages, start=1):
            content = self._normalize_text(page.extract_text() or "")
            if not content:
                continue
            documents.append(
                Document(
                    page_content=content,
                    metadata={
                        "source_path": f"{source}#page={page_number}",
                        "source_file": source,
                        "source_type": "pdf",
                        "title": title,
                        "category": category,
                        "page": page_number,
                        "visibility": "public",
                    },
                )
            )
        return documents

    def _load_docx(self, path: Path) -> Document | None:
        word_document = WordDocument(path)
        blocks = [
            paragraph.text
            for paragraph in word_document.paragraphs
            if paragraph.text.strip()
        ]
        for table in word_document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append(" | ".join(values))
        content = self._normalize_text("\n\n".join(blocks))
        if not content:
            return None

        source = self._relative_source(path)
        internal_names = {
            "Customer Service SOP",
            "Escalation Process",
            "Refund Process",
        }
        return Document(
            page_content=content,
            metadata={
                "source_path": source,
                "source_file": source,
                "source_type": "docx",
                "title": path.stem,
                "category": "Internal Support Process",
                "visibility": "internal" if path.stem in internal_names else "public",
            },
        )


class KnowledgeChunker:
    """Split documents into overlapping chunks sized for local embedding models."""

    def __init__(self, *, chunk_size: int = 1200, chunk_overlap: int = 200) -> None:
        if chunk_size < 200:
            raise ValueError("chunk_size must be at least 200 characters")
        if chunk_overlap < 0 or chunk_overlap >= chunk_size:
            raise ValueError("chunk_overlap must be non-negative and smaller than chunk_size")
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def split_documents(self, documents: list[Document]) -> list[Document]:
        chunks: list[Document] = []
        for document in documents:
            for index, content in enumerate(self._split_text(document.page_content)):
                chunks.append(
                    Document(
                        page_content=content,
                        metadata={**document.metadata, "chunk_index": index},
                    )
                )
        return chunks

    def _split_text(self, text: str) -> list[str]:
        text = text.strip()
        if not text:
            return []
        chunks: list[str] = []
        start = 0
        while start < len(text):
            maximum_end = min(start + self.chunk_size, len(text))
            end = maximum_end
            if maximum_end < len(text):
                search_floor = start + int(self.chunk_size * 0.6)
                candidates = [
                    text.rfind("\n\n", search_floor, maximum_end),
                    text.rfind(". ", search_floor, maximum_end),
                    text.rfind(" ", search_floor, maximum_end),
                ]
                boundary = max(candidates)
                if boundary > start:
                    end = boundary + (1 if text[boundary] == " " else 2)

            content = text[start:end].strip()
            if content:
                chunks.append(content)
            if end >= len(text):
                break
            next_start = max(0, end - self.chunk_overlap)
            if next_start <= start:
                next_start = end
            start = next_start
        return chunks
